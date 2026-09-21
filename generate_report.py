#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
测试报告生成器
根据测试用例xls/xlsx文件和docx模板，自动生成上线后验证报告。
支持两种模板格式：
  - 旧模板：8行5列用例表，基本信息表含执行人员/开始结束时间
  - 新模板：6行4列用例表，基本信息表含编写人员/审批人员，新增需求列表和影响范围表
依赖: pip install openpyxl python-docx
"""

import argparse
import copy
import json
import os
import re
import shutil
import sys
from datetime import datetime

try:
    import openpyxl
except ImportError:
    print("Error: openpyxl is required. Install with: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches
    from docx.text.paragraph import Paragraph
    from lxml import etree
except ImportError:
    print("Error: python-docx is required. Install with: pip install python-docx", file=sys.stderr)
    sys.exit(1)


# ---------------------------------------------------------------------------
# 1. 解析测试用例文件
# ---------------------------------------------------------------------------

def parse_testcase_file(filepath):
    """解析测试用例xls/xlsx文件，返回测试用例列表。"""
    actual_path = filepath
    tmp_file = None

    if filepath.lower().endswith('.xls'):
        with open(filepath, 'rb') as f:
            magic = f.read(4)
        if magic[:2] == b'PK':
            tmp_file = filepath + '.tmp.xlsx'
            shutil.copy2(filepath, tmp_file)
            actual_path = tmp_file
        elif magic[:4] == b'\xd0\xcf\x11\xe0':
            try:
                import xlrd
                return _parse_xls_with_xlrd(filepath)
            except ImportError:
                print("Error: xlrd required for legacy .xls. pip install xlrd", file=sys.stderr)
                sys.exit(1)

    try:
        wb = openpyxl.load_workbook(actual_path)
    finally:
        if tmp_file and os.path.exists(tmp_file):
            os.remove(tmp_file)

    ws = wb.active
    headers = [cell.value for cell in list(ws.iter_rows(min_row=1, max_row=1))[0]]
    col_map = _build_column_map(headers)
    testcases = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        tc = _build_testcase(row, col_map)
        if tc:
            testcases.append(tc)
    return testcases


def _parse_xls_with_xlrd(filepath):
    import xlrd
    wb = xlrd.open_workbook(filepath)
    ws = wb.sheet_by_index(0)
    headers = [ws.cell_value(0, c) for c in range(ws.ncols)]
    col_map = _build_column_map(headers)
    testcases = []
    for row_idx in range(1, ws.nrows):
        row = [ws.cell_value(row_idx, c) for c in range(ws.ncols)]
        tc = _build_testcase(row, col_map)
        if tc:
            testcases.append(tc)
    return testcases


def _build_column_map(headers):
    """依据表头文字自动识别各字段对应的列索引，兼容不同列顺序与列数量。"""
    col_map = {}
    rules = {
        "name": ["用例名称", "用例"],
        "precondition": ["前置条件", "前置"],
        "description": ["描述", "说明"],
        "priority": ["优先级", "优先", "级别"],
        "steps": ["测试步骤", "操作步骤", "步骤", "操作"],
        "expected": ["预期结果", "期望结果", "预期"],
        "case_type": ["用例类型", "类型"],
        "verify_type": ["上线后验证", "验证类型", "验证"],
    }
    for idx, h in enumerate(headers):
        if h is None:
            continue
        h_str = str(h).strip()
        for field, keywords in rules.items():
            for kw in keywords:
                if kw in h_str and field not in col_map:
                    col_map[field] = idx
                    break
    defaults = {"name": 0, "precondition": 1, "description": 2, "priority": 3,
                "steps": 4, "expected": 5, "case_type": 6, "verify_type": 7}
    for field, idx in defaults.items():
        col_map.setdefault(field, idx)
    return col_map


def _build_testcase(row, col_map=None):
    if not row:
        return None
    if col_map is None:
        col_map = {"name": 0, "precondition": 1, "description": 2, "priority": 3,
                   "steps": 4, "expected": 5, "case_type": 6, "verify_type": 7}

    def get_field(field):
        idx = col_map.get(field)
        if idx is None or idx >= len(row):
            return None
        val = row[idx]
        if val is None:
            return None
        s = str(val).strip()
        return s if s else None

    full_name = get_field("name")
    if not full_name:
        return None

    date_tag = ""
    m = re.search(r'[Aa](\d{8})', full_name)
    if m:
        date_tag = m.group(1)

    verify_type = ""
    m = re.search(r'【(项目组验证|地市验证|不具备验证条件)】', full_name)
    if m:
        verify_type = m.group(1)
    else:
        vt = get_field("verify_type")
        if vt:
            verify_type = vt

    clean_name = re.sub(r'【[^】]*】', '', full_name).strip()

    priority_raw = get_field("priority") or "中"
    priority_map = {"高": "高级", "中": "中级", "低": "低级"}

    return {
        "full_name": full_name,
        "name": clean_name,
        "precondition": get_field("precondition") or "",
        "description": get_field("description") or "",
        "priority": priority_map.get(priority_raw, "中级"),
        "steps": get_field("steps") or "",
        "expected": get_field("expected") or "",
        "case_type": "功能用例",
        "verify_type": verify_type or "项目组验证",
        "date_tag": date_tag,
        "status": "通过",
    }


# ---------------------------------------------------------------------------
# 2. 提取项目信息
# ---------------------------------------------------------------------------

def extract_project_info(filepath, testcases):
    basename = os.path.splitext(os.path.basename(filepath))[0]
    basename = re.sub(r'_[a-f0-9]{32}.*$', '', basename)
    # 去掉时间戳后缀 _20260723152354082
    basename = re.sub(r'_\d{17,}$', '', basename)
    parts = [p.strip() for p in basename.split('-') if p.strip()]

    project_name = parts[0] if len(parts) >= 1 else ""
    system_name = parts[1] if len(parts) >= 2 else ""
    feature_name = parts[2] if len(parts) >= 3 else ""
    env_name = parts[3] if len(parts) >= 4 else ""

    date_str = ""
    if testcases and testcases[0].get("date_tag"):
        date_str = testcases[0]["date_tag"]
    else:
        for part in parts:
            m = re.search(r'(\d{8})', part)
            if m:
                date_str = m.group(1)
                cleaned = part.replace(date_str, "").strip()
                if cleaned and "系统" in cleaned:
                    system_name = cleaned
                break
        if not date_str:
            date_str = datetime.now().strftime("%Y%m%d")

    return {
        "project_name": project_name,
        "system_name": system_name,
        "feature_name": feature_name,
        "env_name": env_name,
        "date_str": date_str,
        "report_no": f"自主开发-{system_name}-{date_str}" if system_name else f"测试报告-{date_str}",
    }


# ---------------------------------------------------------------------------
# 3. XML 辅助函数
# ---------------------------------------------------------------------------

def get_paragraph_text(para_elem):
    text = ""
    for run in para_elem.findall(qn('w:r')):
        t = run.find(qn('w:t'))
        if t is not None and t.text:
            text += t.text
    return text.strip()


def set_paragraph_text(para_elem, text):
    """设置段落文本，字体统一为宋体小四"""
    runs = para_elem.findall(qn('w:r'))
    if not runs:
        return
    first_run = runs[0]
    for run in runs[1:]:
        para_elem.remove(run)
    t = first_run.find(qn('w:t'))
    if t is None:
        t = etree.SubElement(first_run, qn('w:t'))
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    rpr = first_run.find(qn('w:rPr'))
    if rpr is None:
        rpr = etree.SubElement(first_run, qn('w:rPr'), attrib={})
        first_run.insert(0, rpr)
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = etree.SubElement(rpr, qn('w:rFonts'))
    rfonts.set(qn('w:ascii'), '宋体')
    rfonts.set(qn('w:eastAsia'), '宋体')
    rfonts.set(qn('w:hAnsi'), '宋体')
    sz = rpr.find(qn('w:sz'))
    if sz is None:
        sz = etree.SubElement(rpr, qn('w:sz'))
    sz.set(qn('w:val'), '24')
    sz_cs = rpr.find(qn('w:szCs'))
    if sz_cs is None:
        sz_cs = etree.SubElement(rpr, qn('w:szCs'))
    sz_cs.set(qn('w:val'), '24')


def set_cell_text_simple(cell, text):
    """设置单元格文本，字体统一为宋体小四，支持\\n换行"""
    from docx.oxml import OxmlElement
    # 处理换行：将文本按\\n分割
    lines = text.split('\n') if '\n' in text else [text]
    
    for para in cell.paragraphs:
        # 清除段落中所有现有的runs
        p_elem = para._element
        for run_elem in list(p_elem.findall(qn('w:r'))):
            p_elem.remove(run_elem)
        
        # 完全用XML方式添加内容，避免add_run()导致顺序错乱
        for i, line in enumerate(lines):
            if i > 0:
                # 添加换行run（只包含<w:br/>）
                br_run = OxmlElement('w:r')
                br_elem = OxmlElement('w:br')
                br_run.append(br_elem)
                p_elem.append(br_run)
            
            # 添加文本run
            run = OxmlElement('w:r')
            rpr = OxmlElement('w:rPr')
            rfonts = OxmlElement('w:rFonts')
            rfonts.set(qn('w:ascii'), '宋体')
            rfonts.set(qn('w:eastAsia'), '宋体')
            rfonts.set(qn('w:hAnsi'), '宋体')
            rpr.append(rfonts)
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), '24')
            rpr.append(sz)
            sz_cs = OxmlElement('w:szCs')
            sz_cs.set(qn('w:val'), '24')
            rpr.append(sz_cs)
            run.append(rpr)
            t = OxmlElement('w:t')
            t.set(qn('xml:space'), 'preserve')
            t.text = line
            run.append(t)
            p_elem.append(run)
        
        return
    
    # 如果没有段落，创建一个
    para = cell.add_paragraph()
    p_elem = para._element
    for i, line in enumerate(lines):
        if i > 0:
            br_run = OxmlElement('w:r')
            br_elem = OxmlElement('w:br')
            br_run.append(br_elem)
            p_elem.append(br_run)
        run = OxmlElement('w:r')
        rpr = OxmlElement('w:rPr')
        rfonts = OxmlElement('w:rFonts')
        rfonts.set(qn('w:ascii'), '宋体')
        rfonts.set(qn('w:eastAsia'), '宋体')
        rfonts.set(qn('w:hAnsi'), '宋体')
        rpr.append(rfonts)
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '24')
        rpr.append(sz)
        sz_cs = OxmlElement('w:szCs')
        sz_cs.set(qn('w:val'), '24')
        rpr.append(sz_cs)
        run.append(rpr)
        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = line
        run.append(t)
        p_elem.append(run)


def _set_run_font_xml(run_elem):
    """通过XML设置run的字体为宋体小四"""
    rpr = run_elem.find(qn('w:rPr'))
    if rpr is None:
        rpr = etree.SubElement(run_elem, qn('w:rPr'), attrib={})
        run_elem.insert(0, rpr)
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = etree.SubElement(rpr, qn('w:rFonts'))
    rfonts.set(qn('w:ascii'), '宋体')
    rfonts.set(qn('w:eastAsia'), '宋体')
    rfonts.set(qn('w:hAnsi'), '宋体')
    sz = rpr.find(qn('w:sz'))
    if sz is None:
        sz = etree.SubElement(rpr, qn('w:sz'))
    sz.set(qn('w:val'), '24')
    sz_cs = rpr.find(qn('w:szCs'))
    if sz_cs is None:
        sz_cs = etree.SubElement(rpr, qn('w:szCs'))
    sz_cs.set(qn('w:val'), '24')


def normalize_all_fonts(doc):
    """统一全文档字体为宋体小四"""
    body = doc.element.body
    for para in body.findall(qn('w:p')):
        for run in para.findall(qn('w:r')):
            rpr = run.find(qn('w:rPr'))
            if rpr is None:
                rpr = etree.SubElement(run, qn('w:rPr'), attrib={})
                run.insert(0, rpr)
            rfonts = rpr.find(qn('w:rFonts'))
            if rfonts is None:
                rfonts = etree.SubElement(rpr, qn('w:rFonts'))
            rfonts.set(qn('w:ascii'), '宋体')
            rfonts.set(qn('w:eastAsia'), '宋体')
            rfonts.set(qn('w:hAnsi'), '宋体')
            sz = rpr.find(qn('w:sz'))
            if sz is None:
                sz = etree.SubElement(rpr, qn('w:sz'))
            sz.set(qn('w:val'), '24')
            sz_cs = rpr.find(qn('w:szCs'))
            if sz_cs is None:
                sz_cs = etree.SubElement(rpr, qn('w:szCs'))
            sz_cs.set(qn('w:val'), '24')
    for tbl in body.findall(qn('w:tbl')):
        for row in tbl.findall(qn('w:tr')):
            for cell in row.findall(qn('w:tc')):
                for para in cell.findall(qn('w:p')):
                    for run in para.findall(qn('w:r')):
                        rpr = run.find(qn('w:rPr'))
                        if rpr is None:
                            rpr = etree.SubElement(run, qn('w:rPr'), attrib={})
                            run.insert(0, rpr)
                        rfonts = rpr.find(qn('w:rFonts'))
                        if rfonts is None:
                            rfonts = etree.SubElement(rpr, qn('w:rFonts'))
                        rfonts.set(qn('w:ascii'), '宋体')
                        rfonts.set(qn('w:eastAsia'), '宋体')
                        rfonts.set(qn('w:hAnsi'), '宋体')
                        sz = rpr.find(qn('w:sz'))
                        if sz is None:
                            sz = etree.SubElement(rpr, qn('w:sz'))
                        sz.set(qn('w:val'), '24')
                        sz_cs = rpr.find(qn('w:szCs'))
                        if sz_cs is None:
                            sz_cs = etree.SubElement(rpr, qn('w:szCs'))
                        sz_cs.set(qn('w:val'), '24')


# ---------------------------------------------------------------------------
# 4. 模板类型检测
# ---------------------------------------------------------------------------

def detect_template(doc):
    """检测模板类型，返回 'new' 或 'old'"""
    if len(doc.tables) < 3:
        return 'old'
    # 新模板特征：Table 1 是 2×5 的需求列表，Table 2 是 3×2 的环境表
    t1 = doc.tables[1]
    if len(t1.rows) == 2 and len(t1.columns) == 5:
        header = t1.rows[0].cells[0].text.strip()
        if '需求编号' in header or '需求名称' in header:
            return 'new'
    return 'old'


# ---------------------------------------------------------------------------
# 5. 填充函数（兼容新旧模板）
# ---------------------------------------------------------------------------

def fill_basic_info(doc, project_info, args):
    """填充基本信息表（Table 0）"""
    table = doc.tables[0]
    template_type = detect_template(doc)

    if template_type == 'new':
        # 新模板: 4×4
        # Row 0: 报告编号 | val(gridSpan=3)
        # Row 1: 编写人员 | val | 编写时间 | val
        # Row 2: 审批人员 | / | 审批时间 | val
        # Row 3: 公司名称 | val(gridSpan=3)
        report_no = args.report_no or project_info['report_no']
        executor = args.executor or ""
        company = args.company or ""
        today = datetime.now().strftime("%Y-%m-%d")

        # Row 0: 报告编号
        set_cell_text_simple(table.rows[0].cells[1], report_no)

        # Row 1: 编写人员 + 编写时间
        set_cell_text_simple(table.rows[1].cells[1], executor)
        set_cell_text_simple(table.rows[1].cells[3], today)

        # Row 2: 审批人员 + 审批时间 (保持默认)

        # Row 3: 公司名称
        set_cell_text_simple(table.rows[3].cells[1], company)
    else:
        # 旧模板: 4×4
        # Row 0: 报告编号 | val | 执行人员 | val
        # Row 1: 编写日期 | val | 开始时间 | 结束时间
        # Row 2: 公司名称 | val | (空) | (空)
        report_no = args.report_no or project_info['report_no']
        executor = args.executor or ""
        company = args.company or ""
        start_time = args.start_time or ""
        end_time = args.end_time or ""
        today = datetime.now().strftime("%Y-%m-%d")

        set_cell_text_simple(table.rows[0].cells[1], report_no)
        set_cell_text_simple(table.rows[0].cells[3], executor)
        set_cell_text_simple(table.rows[1].cells[1], today)
        set_cell_text_simple(table.rows[1].cells[3], f"{start_time}~{end_time}" if start_time else "")
        set_cell_text_simple(table.rows[2].cells[1], company)


def fill_requirement_table(doc, project_info, testcases):
    """填充需求列表（新模板 Table 1）"""
    if detect_template(doc) != 'new':
        return
    table = doc.tables[1]
    if len(table.rows) < 2:
        return
    # Row 1: 需求编号 | 需求名称 | 需求说明 | 测试人员 | 完成时间
    today = datetime.now().strftime("%Y-%m-%d")
    feature = project_info['feature_name'] or project_info['system_name'] or "需求"
    executor = ""  # 从args获取需要传参，这里简化

    cells = table.rows[1].cells
    set_cell_text_simple(cells[0], "1")
    set_cell_text_simple(cells[1], feature)
    set_cell_text_simple(cells[2], feature)
    set_cell_text_simple(cells[3], "")
    set_cell_text_simple(cells[4], today)


def fill_web_env(doc, args):
    """填充测试环境表"""
    template_type = detect_template(doc)

    if template_type == 'new':
        # 新模板 Table 2: 3×2
        # Row 0: 测试地址URL | url
        # Row 1: 测试浏览器 | browser
        # Row 2: 测试账号/密码 | account
        table = doc.tables[2]
        browser = args.browser or "Google Chrome"
        account = args.account or ""
        set_cell_text_simple(table.rows[1].cells[1], browser)
        set_cell_text_simple(table.rows[2].cells[1], account)
    else:
        # 旧模板 Table 1: 2×2
        # Row 0: 测试浏览器 | browser
        # Row 1: 测试账号 | account
        table = doc.tables[1]
        browser = args.browser or "Google Chrome"
        account = args.account or ""
        set_cell_text_simple(table.rows[0].cells[1], browser)
        set_cell_text_simple(table.rows[1].cells[1], account)


def fill_impact_table(doc, project_info):
    """填充上线功能影响范围表（新模板 Table 6）"""
    if detect_template(doc) != 'new':
        return
    # 找到影响范围表（最后一个2×5的表）
    for table in doc.tables:
        if len(table.rows) == 2 and len(table.columns) == 5:
            header = table.rows[0].cells[0].text.strip()
            if '需求单号' in header or '上线内容' in header:
                feature = project_info['feature_name'] or project_info['system_name'] or ""
                cells = table.rows[1].cells
                set_cell_text_simple(cells[0], "1")
                set_cell_text_simple(cells[1], feature)
                set_cell_text_simple(cells[2], feature)
                set_cell_text_simple(cells[3], feature)
                set_cell_text_simple(cells[4], "")
                return


def fill_testcase_table(table, tc, index, template_type):
    """填充单个测试用例表格"""
    if template_type == 'new':
        # 新模板: 6×4
        # Row 0: 用例标题 | name | 测试状态 | 通过
        # Row 1: 前置条件 | condition | 用例类型 | 功能用例
        # Row 2: 账号/数据 | account | 优先等级 | level
        # Row 3: 测试步骤 | steps(gridSpan=3)
        # Row 4: 预期结果 | expected(gridSpan=3)
        # Row 5: 实际结果 | actual(gridSpan=3)
        set_cell_text_simple(table.rows[0].cells[1], tc['name'])
        set_cell_text_simple(table.rows[0].cells[3], tc['status'])
        set_cell_text_simple(table.rows[1].cells[1], tc['precondition'])
        set_cell_text_simple(table.rows[1].cells[3], tc['case_type'])
        set_cell_text_simple(table.rows[2].cells[1], "同环境信息")
        set_cell_text_simple(table.rows[2].cells[3], tc['priority'])
        set_cell_text_simple(table.rows[3].cells[1], tc['steps'])
        set_cell_text_simple(table.rows[4].cells[1], tc['expected'])
        set_cell_text_simple(table.rows[5].cells[1], tc['expected'])
    else:
        # 旧模板: 8×5
        # Row 0: 用例标题 | name | 测试状态 | 通过
        # Row 1: 前置条件 | condition | 用例类型 | 功能用例
        # Row 2: 账号/数据 | 同环境信息 | 优先等级 | level
        # Row 3-5: 测试步骤(合并3行)
        # Row 6: 预期结果(合并)
        # Row 7: 实际结果(合并)
        set_cell_text_simple(table.rows[0].cells[1], tc['name'])
        set_cell_text_simple(table.rows[0].cells[3], tc['status'])
        set_cell_text_simple(table.rows[1].cells[1], tc['precondition'])
        set_cell_text_simple(table.rows[1].cells[3], tc['case_type'])
        set_cell_text_simple(table.rows[2].cells[1], "同环境信息")
        set_cell_text_simple(table.rows[2].cells[3], tc['priority'])
        set_cell_text_simple(table.rows[3].cells[1], tc['steps'])
        set_cell_text_simple(table.rows[6].cells[1], tc['expected'])
        set_cell_text_simple(table.rows[7].cells[1], tc['expected'])


def generate_testcase_elements(doc, testcases, template_type):
    """为每个测试用例生成 Heading3 + 表格 + 截图段落"""
    body = doc.element.body
    children = list(body)

    # 找到关键段落的索引
    results_idx = None  # "测试执行结果"
    summary_idx = None  # "测试总结" 或 "上线功能影响范围"

    for i, child in enumerate(children):
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag == 'p':
            text = get_paragraph_text(child)
            if '测试执行结果' in text or '验证执行结果' in text:
                results_idx = i
            if results_idx is not None and summary_idx is None:
                if '测试总结' in text or '上线功能影响范围' in text:
                    summary_idx = i

    if results_idx is None or summary_idx is None:
        return

    # 找到第一个测试用例块作为模板（results_idx 之后的第一个 heading + table + screenshot）
    template_heading_elem = None
    template_table_elem = None
    template_screenshot_elem = None

    for i in range(results_idx + 1, summary_idx):
        child = children[i]
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        text = get_paragraph_text(child) if tag == 'p' else ''

        if template_heading_elem is None and tag == 'p' and ('测试用例' in text or '用例' in text):
            template_heading_elem = child
        elif template_heading_elem is not None and template_table_elem is None and tag == 'tbl':
            template_table_elem = child
        elif template_table_elem is not None and template_screenshot_elem is None and tag == 'p' and '测试截图' in text:
            template_screenshot_elem = child
            break

    if template_heading_elem is None or template_table_elem is None:
        return

    # 删除 results_idx+1 到 summary_idx-1 之间的所有元素
    to_remove = children[results_idx + 1:summary_idx]
    for elem in to_remove:
        body.remove(elem)

    # 重新获取 children 列表（因为删除了元素）
    children = list(body)
    # 找到 summary 元素的新位置
    insert_before = None
    for child in children:
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag == 'p':
            text = get_paragraph_text(child)
            if '测试总结' in text or '上线功能影响范围' in text:
                insert_before = child
                break

    # 为每个测试用例生成元素
    for idx, tc in enumerate(testcases):
        # Heading
        heading = copy.deepcopy(template_heading_elem)
        set_paragraph_text(heading, f"测试用例{idx + 1}：{tc['name']}")
        if insert_before is not None:
            body.insert(body.index(insert_before), heading)
        else:
            body.append(heading)

        # Table - 用XML方式填充
        table_elem = copy.deepcopy(template_table_elem)
        _fill_table_xml(table_elem, tc, template_type)
        if insert_before is not None:
            body.insert(body.index(insert_before), table_elem)
        else:
            body.append(table_elem)

        # 截图段落
        if template_screenshot_elem is not None:
            screenshot = copy.deepcopy(template_screenshot_elem)
            set_paragraph_text(screenshot, f"测试截图：{tc['name']}")
            if insert_before is not None:
                body.insert(body.index(insert_before), screenshot)
            else:
                body.append(screenshot)


def _fill_table_xml(table_elem, tc, template_type):
    """直接用XML方式填充表格内容"""
    rows = table_elem.findall(qn('w:tr'))

    if template_type == 'new' and len(rows) >= 6:
        # 6×4 新模板
        _set_row_cell_text(rows[0], 1, tc['name'])
        _set_row_cell_text(rows[0], 3, tc['status'])
        _set_row_cell_text(rows[1], 1, tc['precondition'])
        _set_row_cell_text(rows[1], 3, tc['case_type'])
        _set_row_cell_text(rows[2], 1, "同环境信息")
        _set_row_cell_text(rows[2], 3, tc['priority'])
        _set_row_cell_text(rows[3], 1, tc['steps'])
        _set_row_cell_text(rows[4], 1, tc['expected'])
        _set_row_cell_text(rows[5], 1, tc['expected'])
    elif len(rows) >= 8:
        # 8×5 旧模板
        _set_row_cell_text(rows[0], 1, tc['name'])
        _set_row_cell_text(rows[0], 3, tc['status'])
        _set_row_cell_text(rows[1], 1, tc['precondition'])
        _set_row_cell_text(rows[1], 3, tc['case_type'])
        _set_row_cell_text(rows[2], 1, "同环境信息")
        _set_row_cell_text(rows[2], 3, tc['priority'])
        _set_row_cell_text(rows[3], 1, tc['steps'])
        _set_row_cell_text(rows[6], 1, tc['expected'])
        _set_row_cell_text(rows[7], 1, tc['expected'])


def _set_row_cell_text(row_elem, col_idx, text):
    """设置行中指定列的文本，支持\\n换行"""
    cells = row_elem.findall(qn('w:tc'))
    if col_idx >= len(cells):
        return
    cell = cells[col_idx]
    paras = cell.findall(qn('w:p'))
    if not paras:
        return
    para = paras[0]
    # 清除段落中所有现有的runs
    for run_elem in list(para.findall(qn('w:r'))):
        para.remove(run_elem)
    
    # 处理换行：将文本按\n分割
    lines = text.split('\n') if '\n' in text else [text]
    
    for i, line in enumerate(lines):
        if i > 0:
            # 添加换行run（只包含<w:br/>）
            br_run = etree.SubElement(para, qn('w:r'))
            br_elem = etree.SubElement(br_run, qn('w:br'))
        
        # 添加文本run
        run = etree.SubElement(para, qn('w:r'))
        rpr = etree.SubElement(run, qn('w:rPr'))
        rfonts = etree.SubElement(rpr, qn('w:rFonts'))
        rfonts.set(qn('w:ascii'), '宋体')
        rfonts.set(qn('w:eastAsia'), '宋体')
        rfonts.set(qn('w:hAnsi'), '宋体')
        sz = etree.SubElement(rpr, qn('w:sz'))
        sz.set(qn('w:val'), '24')
        sz_cs = etree.SubElement(rpr, qn('w:szCs'))
        sz_cs.set(qn('w:val'), '24')
        t = etree.SubElement(run, qn('w:t'))
        t.set(qn('xml:space'), 'preserve')
        t.text = line


def update_summary(doc, project_info, testcases, template_type):
    """更新测试总结段落"""
    date_str = project_info['date_str']
    feature = project_info['feature_name'] or project_info['system_name'] or "需求"
    total = len(testcases)

    # 统计验证类型
    verify_counts = {}
    for tc in testcases:
        vt = tc.get('verify_type', '项目组验证')
        verify_counts[vt] = verify_counts.get(vt, 0) + 1

    passed = sum(1 for tc in testcases if tc.get('status') == '通过')

    for para in doc.paragraphs:
        if '总计需求' in para.text or '总计需求' in para.text:
            if template_type == 'new':
                summary = f"{feature}，总计需求1条。执行测试用例{total}条，测试通过{passed}条。"
            else:
                parts = []
                for vt, count in verify_counts.items():
                    parts.append(f"{vt}用例{count}个已全部验证通过")
                summary_text = "。".join(parts) if parts else f"共{total}条测试用例已全部验证通过"
                summary = f"{date_str}上线版本，总计需求1个,共{total}条测试用例.{summary_text}。"
            set_paragraph_text(para._p, summary)
            break


# ---------------------------------------------------------------------------
# 6. 主流程
# ---------------------------------------------------------------------------

def _find_image_file(screenshots_dir, seq):
    """在截图目录中查找指定序号(seq)的图片文件"""
    extensions = ['.jpg', '.jpeg', '.png', '.bmp', '.gif', '.webp']
    for ext in extensions:
        path = os.path.join(screenshots_dir, f"{seq}{ext}")
        if os.path.isfile(path):
            return path
    return None


def insert_screenshots(doc, testcases, screenshots_dir):
    """在每个测试用例的"测试截图"段落后批量插入对应用例序号的照片"""
    from docx.oxml.ns import qn as _qn
    body = doc.element.body
    # 收集所有"测试截图："段落元素（按出现顺序）
    screenshot_paras = []
    for child in body:
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag == 'p':
            text = get_paragraph_text(child)
            if '测试截图' in text:
                screenshot_paras.append(child)

    inserted = 0
    for idx, para_elem in enumerate(screenshot_paras):
        seq = idx + 1  # 用例序号从1开始，与测试用例索引对应
        image_path = _find_image_file(screenshots_dir, seq)
        if image_path is None:
            continue
        try:
            # 用doc.add_paragraph创建（python-docx oxml元素，支持add_run/add_picture），再移动到截图段落后
            new_para = doc.add_paragraph()
            run = new_para.add_run()
            run.add_picture(image_path, width=Cm(14))
            para_elem.addnext(new_para._p)
            inserted += 1
        except Exception as e:
            print(f"插入截图 {seq} 失败: {e}", file=sys.stderr)

    print(json.dumps({"status": "success", "inserted_screenshots": inserted}, ensure_ascii=False))


def generate_report(testcase_file, template_file, output_file, args):
    testcases = parse_testcase_file(testcase_file)
    if not testcases:
        print(json.dumps({"status": "error", "message": "未找到测试用例"}, ensure_ascii=False))
        sys.exit(1)

    project_info = extract_project_info(testcase_file, testcases)
    template_type = detect_template(Document(template_file))

    doc = Document(template_file)

    # 填充基本信息
    fill_basic_info(doc, project_info, args)

    # 填充需求列表（新模板）
    fill_requirement_table(doc, project_info, testcases)

    # 填充环境信息
    fill_web_env(doc, args)

    # 填充影响范围表（新模板）
    fill_impact_table(doc, project_info)

    # 生成测试用例条目
    generate_testcase_elements(doc, testcases, template_type)

    # 批量插入截图
    screenshots_dir = getattr(args, 'screenshots_dir', '')
    if screenshots_dir and os.path.isdir(screenshots_dir):
        insert_screenshots(doc, testcases, screenshots_dir)

    # 更新测试总结
    update_summary(doc, project_info, testcases, template_type)

    # 统一字体
    normalize_all_fonts(doc)

    doc.save(output_file)

    result = {
        "status": "success",
        "output_file": os.path.abspath(output_file),
        "testcase_count": len(testcases),
        "template_type": template_type,
        "project_info": project_info,
    }
    print(json.dumps(result, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(description='测试报告生成器')
    parser.add_argument('--testcase', required=True, help='测试用例xls/xlsx文件路径')
    parser.add_argument('--template', required=True, help='报告模板docx文件路径')
    parser.add_argument('--output', required=True, help='输出报告docx文件路径')
    parser.add_argument('--report-no', default='', help='报告编号')
    parser.add_argument('--executor', default='', help='执行人员/编写人员')
    parser.add_argument('--company', default='', help='公司名称')
    parser.add_argument('--browser', default='', help='测试浏览器')
    parser.add_argument('--account', default='', help='测试账号')
    parser.add_argument('--start-time', default='', help='开始时间')
    parser.add_argument('--end-time', default='', help='结束时间')
    parser.add_argument('--screenshots-dir', default='', help='截图文件夹路径，图片按用例序号命名(1.jpg, 2.png...)')

    args = parser.parse_args()

    if not os.path.exists(args.testcase):
        print(json.dumps({"status": "error", "message": f"测试用例文件不存在: {args.testcase}"}, ensure_ascii=False))
        sys.exit(1)

    if not os.path.exists(args.template):
        print(json.dumps({"status": "error", "message": f"模板文件不存在: {args.template}"}, ensure_ascii=False))
        sys.exit(1)

    generate_report(args.testcase, args.template, args.output, args)


if __name__ == "__main__":
    main()
